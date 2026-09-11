import json
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def main():
    data = json.loads(sys.stdin.read())
    result = data["result"]
    method_name = data.get("method_name", "Método numérico")

    wb = Workbook()
    ws = wb.active
    ws.title = "Iteraciones"

    title = f"{method_name} — {result.get('function', '')}"
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(result.get("columns", []))))

    columns = result.get("columns", [])
    for col, name in enumerate(columns, start=1):
        cell = ws.cell(row=3, column=col, value=name)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4F46E5")
        cell.alignment = Alignment(horizontal="center")

    for r, row in enumerate(result.get("table", []), start=4):
        for c, name in enumerate(columns, start=1):
            value = row.get(name)
            ws.cell(row=r, column=c, value=value)

    thin = Side(style="thin", color="D1D5DB")
    for row in ws.iter_rows(min_row=3, max_row=3+len(result.get("table", [])), min_col=1, max_col=max(1, len(columns))):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(horizontal="center")

    for col in range(1, max(1, len(columns))+1):
        letter = get_column_letter(col)
        ws.column_dimensions[letter].width = min(24, max(12, len(str(ws.cell(3, col).value or "")) + 4))

    ws.freeze_panes = "A4"

    summary = wb.create_sheet("Resumen")
    summary.append(["Campo", "Valor"])
    summary["A1"].font = Font(bold=True)
    summary["B1"].font = Font(bold=True)
    summary.append(["Estado", "Convergencia alcanzada" if result.get("converged") else "No convergió"])
    summary.append(["Mensaje", result.get("message")])
    summary.append(["x final", result.get("final_x")])
    summary.append(["f(x)", result.get("final_fx")])
    summary.append(["Iteraciones", result.get("iterations")])
    if result.get("final_y") is not None:
        summary.append(["y final", result.get("final_y")])

    out = data["output_path"]
    wb.save(out)
    print(json.dumps({"ok": True, "path": out}))

if __name__ == "__main__":
    main()
